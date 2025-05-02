import tkinter as tk
from tkinter import *
from tkinter import messagebox
from tkinter import ttk
from openpyxl import load_workbook

window = tk.Tk()

height = 200 
width = 330 
x = (window.winfo_screenwidth()//2)-(width//2) 
y = (window.winfo_screenheight()//2)-(height//1) 
window.geometry('{}x{}+{}+{}'.format(width, height, x, y))

window.title("STUDENT SCORE TRACKER")
window.config(bg="#EFEEEA")
img = PhotoImage(file='J:\\ITCS103_DE VILLA_1A\\quiz_2\\myicon.png')
window.iconphoto(False,img)

frame = tk.Frame(bg="#EFEEEA")
s_frame = tk.Frame(bg="#EFEEEA")

def validate_input():
    name = name_entry.get()
    grade = grade_entry.get()

    if not name or not grade:
        messagebox.showerror("Error", "All Fields are Required")
        return False

    try:
        int(grade)
    except ValueError:
        messagebox.showerror("Error", "Grade must be a number")
        return False

    return True

def save_to_excel():
    if not validate_input():
        return

    name = name_entry.get()
    grade = int(grade_entry.get())

    
    if grade <= 74:
        remark = "Failed"
    else:
        remark = "Passed"

    wb = load_workbook("student_scores.xlsx")
    ws = wb["student_scores"]
    ws.append([name, grade, remark])
    wb.save("student_scores.xlsx")
    messagebox.showinfo("Success", "Data Saved Successfully")



tk.Label(frame, text="STUDENT SCORE TRACKER", font=('Helvetica', 12),bg="#EFEEEA").grid(row=0, column=0, columnspan=2, pady=20)
tk.Label(frame, text="Name",bg="#EFEEEA").grid(row=1, column=0)
tk.Label(frame, text="Enter your Grade",bg="#EFEEEA").grid(row=2, column=0)

name_entry = tk.Entry(frame)
grade_entry = tk.Entry(frame)

name_entry.grid(row=1, column=1)
grade_entry.grid(row=2, column=1)


save_button = tk.Button(frame, text="Save", command=save_to_excel, padx=10,bg="#273F4F",fg="white")
save_button.grid(row=3, column=1, columnspan=2, pady=10)


def open_view_window():
    view_window = tk.Toplevel(window)
    view_window.title("Student Entries")

    # view_window.geometry("350x300")

    height = 350 
    width = 300 
    x = (view_window.winfo_screenwidth()//2)-(width//1) 
    y = (view_window.winfo_screenheight()//2)-(height//2) 
    view_window.geometry('{}x{}+{}+{}'.format(width, height, x, y))


    view_window.config(bg="#EFEEEA")
    img = PhotoImage(file='J:\\ITCS103_DE VILLA_1A\\quiz_2\\myicon.png')
    view_window.iconphoto(False,img)

    columns = ("Name", "Grade", "Remark")
    tree = ttk.Treeview(view_window, columns=columns, show="headings")

    for col in columns:
        tree.heading(col, text=col)
        tree.column(col, width=100)

    tree.pack(fill="both", expand=True, padx=10, pady=10)
#average
    total = 0
    count = 0

    wb = load_workbook("student_scores.xlsx")
    ws = wb["student_scores"]

    for row in ws.iter_rows(min_row=2, values_only=True):
        tree.insert("", "end", values=row)
        try:
            grade = int(row[1])
            total += grade
            count += 1
        except:
            continue

    if count > 0:
        average = total / count
        avg_label = tk.Label(view_window, text=f"Average Grade: {average:.2f}", font=('Helvetica', 10, 'bold'))
        avg_label.pack(pady=10)



view_button = tk.Button(frame, text="View Data", command=open_view_window,bg="#169976",padx=5)
view_button.grid(row=3, column=0,columnspan=1)

frame.pack(expand=True)
s_frame.pack(expand=True)

window.mainloop()
